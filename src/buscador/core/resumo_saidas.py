# -*- coding: utf-8 -*-
"""
Funções que "olham" a pasta saidas/ (onde o motor grava planilhas, logs e
os arquivos baixados) para o aplicativo com tela (dashboard) mostrar
números e listas -- sem nunca mudar nada lá dentro. Só leitura.

Ficam aqui (e não dentro das telas) por dois motivos:
1. dá para testar com pytest, sem abrir tela nenhuma;
2. mais de uma tela usa a mesma coisa (Início e Mapear mostram planilhas).

Cuidado de desempenho: saidas/telegram/ tem milhares de arquivos grandes
(às vezes num HD externo). Por isso usamos os.scandir -- ele lista a pasta
sem abrir os arquivos -- e nunca lemos o conteúdo de um arquivo baixado.
"""
import datetime
import os
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from buscador.core.verificacao_links import COR

# Pasta saidas/ na raiz do projeto (sobe de src/buscador/core/ até a raiz).
# As outras partes do app leem ESTA variável na hora de usar (não copiam o
# valor), então os testes podem trocá-la por uma pasta temporária.
SAIDAS = Path(__file__).resolve().parent.parent.parent.parent / "saidas"

# Pastas que a busca por planilhas NÃO abre: "arquivos" (downloads do
# Telegram, milhares de PDFs) e "repetidos" (cópias separadas pelo Samuel).
# Entrar nelas só deixaria tudo lento, e lá não tem planilha nossa.
PASTAS_IGNORADAS = {"arquivos", "repetidos", "__pycache__"}

# Cor de fundo da linha na planilha -> bolinha que a tela mostra.
# As cores vêm do mesmo lugar que pinta a planilha (verificacao_links.COR),
# para as duas coisas nunca ficarem diferentes.
EMOJI_POR_COR = {
    COR["verde"].fgColor.rgb[-6:]: "🟢",
    COR["requer login"].fgColor.rgb[-6:]: "🔵",
    COR["quebrado"].fgColor.rgb[-6:]: "🔴",
}
LEGENDA_STATUS = {
    "🟢": "baixa direto (vivo + PDF)",
    "🔵": "precisa de login",
    "🔴": "link quebrado",
    "⚪": "só vivo (não deu para ver o PDF)",
}


@dataclass
class Planilha:
    """Uma planilha .xlsx encontrada em saidas/ -- só o que a tela precisa."""
    caminho: Path
    modificada_em: datetime.datetime
    tamanho_bytes: int


def pasta_saidas() -> Path:
    """Devolve a pasta saidas/ atual (lida na hora, ver comentário de SAIDAS)."""
    return SAIDAS


def listar_planilhas(pasta: Path | None = None, profundidade_maxima: int = 3) -> list[Planilha]:
    """Procura arquivos .xlsx dentro de saidas/ (e subpastas, até
    "profundidade_maxima" níveis), da mais nova para a mais velha.
    Ignora arquivos temporários do Excel ("~$nome.xlsx", que aparecem
    enquanto a planilha está aberta)."""
    pasta = Path(pasta) if pasta else pasta_saidas()
    encontradas = []
    _procurar_xlsx(pasta, profundidade_maxima, encontradas)
    encontradas.sort(key=lambda p: p.modificada_em, reverse=True)
    return encontradas


def _procurar_xlsx(pasta: Path, profundidade_restante: int, encontradas: list) -> None:
    """Função "recursiva" (que chama a si mesma) -- desce pasta por pasta,
    até acabar a profundidade permitida."""
    try:
        entradas = list(os.scandir(pasta))
    except OSError:
        return  # pasta não existe ou não deu para abrir: só não acha nada
    for entrada in entradas:
        try:
            if entrada.is_dir():
                if profundidade_restante > 0 and entrada.name not in PASTAS_IGNORADAS:
                    _procurar_xlsx(Path(entrada.path), profundidade_restante - 1, encontradas)
            elif entrada.name.lower().endswith(".xlsx") and not entrada.name.startswith("~$"):
                info = entrada.stat()
                encontradas.append(Planilha(
                    caminho=Path(entrada.path),
                    modificada_em=datetime.datetime.fromtimestamp(info.st_mtime),
                    tamanho_bytes=info.st_size,
                ))
        except OSError:
            continue  # um arquivo com problema não derruba a lista inteira


def contar_arquivos_telegram(pasta: Path | None = None) -> tuple[int, int]:
    """Conta os arquivos já baixados do Telegram e soma o tamanho deles.
    Devolve (quantidade, total_em_bytes). Olha só dentro das pastas
    chamadas "arquivos" debaixo de saidas/telegram/ (é onde o download
    grava) -- inclusive a subpasta "repetidos". Arquivos ".tmp" são
    downloads pela metade, e não contam."""
    pasta = Path(pasta) if pasta else pasta_saidas() / "telegram"
    quantidade, total = 0, 0
    for pasta_arquivos in _achar_pastas_arquivos(pasta, profundidade_restante=3):
        q, t = _contar_recursivo(pasta_arquivos)
        quantidade += q
        total += t
    return quantidade, total


def _achar_pastas_arquivos(pasta: Path, profundidade_restante: int) -> list[Path]:
    """Acha as pastas chamadas "arquivos" (sem entrar nelas para procurar)."""
    achadas = []
    try:
        entradas = list(os.scandir(pasta))
    except OSError:
        return achadas
    for entrada in entradas:
        try:
            if not entrada.is_dir():
                continue
            if entrada.name == "arquivos":
                achadas.append(Path(entrada.path))
            elif profundidade_restante > 0:
                achadas.extend(_achar_pastas_arquivos(Path(entrada.path), profundidade_restante - 1))
        except OSError:
            continue
    return achadas


def _contar_recursivo(pasta: Path) -> tuple[int, int]:
    quantidade, total = 0, 0
    try:
        entradas = list(os.scandir(pasta))
    except OSError:
        return 0, 0
    for entrada in entradas:
        try:
            if entrada.is_dir():
                q, t = _contar_recursivo(Path(entrada.path))
                quantidade += q
                total += t
            elif not entrada.name.endswith(".tmp"):
                quantidade += 1
                total += entrada.stat().st_size
        except OSError:
            continue
    return quantidade, total


def formatar_tamanho(total_bytes: int) -> str:
    """1234567890 -> "1,1 GB" (com vírgula, do jeito brasileiro)."""
    valor = float(total_bytes)
    for unidade in ("bytes", "KB", "MB", "GB"):
        if valor < 1024 or unidade == "GB":
            if unidade == "bytes":
                return f"{int(valor)} bytes"
            return f"{valor:.1f} {unidade}".replace(".", ",")
        valor /= 1024
    return f"{valor:.1f} TB".replace(".", ",")  # nunca chega aqui; só para garantir


def ler_previa_planilha(caminho: Path, limite: int = 200) -> tuple[list[str], list[list], list[str]]:
    """Lê só o comecinho de uma planilha (até "limite" linhas, sem contar o
    cabeçalho) para mostrar na tela. Devolve (colunas, linhas, bolinhas):
    "bolinhas" tem uma bolinha de status por linha (🟢🔵🔴⚪), descoberta
    pela COR de fundo que o motor pintou na linha.

    Usa read_only=True: o openpyxl lê a planilha aos pouquinhos em vez de
    carregar tudo na memória -- importante, porque há planilhas com 46 mil
    linhas em saidas/."""
    wb = load_workbook(caminho, read_only=True)
    try:
        ws = wb.worksheets[0]  # a primeira aba ("Links", nas planilhas do motor)
        colunas, linhas, bolinhas = [], [], []
        for numero, linha in enumerate(ws.iter_rows()):
            if numero == 0:
                colunas = [str(c.value) if c.value is not None else f"coluna_{i + 1}"
                           for i, c in enumerate(linha)]
                continue
            if numero > limite:
                break
            linhas.append([c.value for c in linha])
            bolinhas.append(_bolinha_da_celula(linha[0] if linha else None))
        return colunas, linhas, bolinhas
    finally:
        wb.close()  # em read_only o arquivo fica aberto até fechar -- no Windows isso travaria o arquivo


def _bolinha_da_celula(celula) -> str:
    try:
        cor = celula.fill.fgColor.rgb  # ex.: "00C6EFCE"
    except AttributeError:
        return "⚪"
    if not isinstance(cor, str):
        return "⚪"
    return EMOJI_POR_COR.get(cor[-6:].upper(), "⚪")
