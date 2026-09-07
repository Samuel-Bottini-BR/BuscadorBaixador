# -*- coding: utf-8 -*-
from openpyxl import load_workbook

from buscador.adapters.base import Item
from buscador.core.planilha import COLUNAS, OPCOES_AVALIACAO, gerar_planilha


def _itens_de_exemplo():
    return [
        Item(titulo_original="Obra verde", link="https://exemplo.com/a.pdf",
             titulo_pt="Green work", autor="Fulano", ano="1999", tipo="pdf direto",
             status_link="vivo (pdf direto)", explicacao="achado no forum",
             fonte="Grand Sud Médiéval", extra={"categoria": "verde"}),
        Item(titulo_original="Obra login", link="https://academia.edu/b",
             tipo="precisa login", status_link="requer login",
             fonte="Grand Sud Médiéval", extra={"categoria": "requer login"}),
        Item(titulo_original="Obra quebrada", link="https://exemplo.com/sumiu",
             tipo="quebrado", status_link="quebrado",
             fonte="Grand Sud Médiéval", extra={"categoria": "quebrado"}),
        Item(titulo_original="Obra branca", link="https://exemplo.com/pagina",
             tipo="página", status_link="vivo (sem pdf)",
             fonte="Grand Sud Médiéval", extra={"categoria": "branco"}),
    ]


def test_cabecalho_e_valores(tmp_path):
    destino = gerar_planilha(_itens_de_exemplo(), tmp_path / "saida.xlsx")
    wb = load_workbook(destino)
    ws = wb["Links"]
    assert [c.value for c in ws[1]] == COLUNAS
    assert ws.cell(row=2, column=COLUNAS.index("titulo_original") + 1).value == "Obra verde"
    assert ws.cell(row=2, column=COLUNAS.index("titulo_pt") + 1).value == "Green work"
    assert ws.max_row == 5  # cabecalho + 4 itens


def test_linha_pintada_pela_categoria(tmp_path):
    destino = gerar_planilha(_itens_de_exemplo(), tmp_path / "saida.xlsx")
    ws = load_workbook(destino)["Links"]
    # linha 2 = verde, linha 3 = requer login (azul), linha 4 = quebrado (vermelho)
    assert ws.cell(row=2, column=1).fill.fgColor.rgb == "00C6EFCE"
    assert ws.cell(row=3, column=1).fill.fgColor.rgb == "00BDD7EE"
    assert ws.cell(row=4, column=1).fill.fgColor.rgb == "00FFC7CE"


def test_linha_branca_nao_e_pintada(tmp_path):
    destino = gerar_planilha(_itens_de_exemplo(), tmp_path / "saida.xlsx")
    ws = load_workbook(destino)["Links"]
    assert ws.cell(row=5, column=1).fill.patternType is None


def test_dropdown_de_avaliacao_existe_com_as_opcoes_certas(tmp_path):
    destino = gerar_planilha(_itens_de_exemplo(), tmp_path / "saida.xlsx")
    ws = load_workbook(destino)["Links"]
    validacoes = list(ws.data_validations.dataValidation)
    assert len(validacoes) == 1
    formula = validacoes[0].formula1
    for opcao in OPCOES_AVALIACAO:
        assert opcao in formula


def test_planilha_vazia_nao_quebra(tmp_path):
    destino = gerar_planilha([], tmp_path / "vazia.xlsx")
    ws = load_workbook(destino)["Links"]
    assert ws.max_row == 1  # so o cabecalho


def test_larguras_de_coluna_aplicadas(tmp_path):
    destino = gerar_planilha(_itens_de_exemplo(), tmp_path / "saida.xlsx")
    ws = load_workbook(destino)["Links"]
    assert ws.column_dimensions["A"].width == 40  # titulo_original
    assert ws.column_dimensions["E"].width == 50  # link
