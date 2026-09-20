# -*- coding: utf-8 -*-
import pytest
from openpyxl import load_workbook

from buscador import cli
from buscador.adapters.base import Item
from buscador.adapters.gallica import GallicaAdapter
from buscador.adapters.internet_archive import InternetArchiveAdapter
from buscador.core.acao_humana import TIPO_CHAVE_API, AcaoHumanaNecessaria


class AdapterFake:
    def __init__(self, url):
        self.url = url

    def iter_itens(self):
        yield Item(titulo_original="Obra Teste", link="https://exemplo.com/a.pdf",
                    autor="Fulano", fonte="Fonte Fake", extra={"idioma_origem": "fr"})
        yield Item(titulo_original="Outra Obra", link="https://academia.edu/b",
                    autor="Ciclano", fonte="Fonte Fake", extra={"idioma_origem": "fr"})


class AdapterFakePrecisaChave:
    def __init__(self, url):
        self.url = url

    def iter_itens(self):
        raise AcaoHumanaNecessaria(TIPO_CHAVE_API, "cadastre-se em tal-site e cole a chave", site="tal-site")


def test_escolher_adapter_por_dominio():
    url = "https://grand-sud-medieval.fr/forum/viewtopic.php?f=14&t=9264"
    assert cli.escolher_adapter(url) == "phpbb"


def test_escolher_adapter_gallica_por_dominio():
    url = 'https://gallica.bnf.fr/SRU?operation=searchRetrieve&version=1.2&query=gallica all "Clavius"'
    assert cli.escolher_adapter(url) == "gallica"


def test_escolher_adapter_internet_archive_por_dominio():
    url = "https://archive.org/details/algumitem"
    assert cli.escolher_adapter(url) == "internet_archive"


def test_construir_adapter_internet_archive():
    adapter = cli.construir_adapter("internet_archive", "subject:theology")
    assert isinstance(adapter, InternetArchiveAdapter)
    assert adapter.consulta == "subject:theology"


def test_extrair_consulta_gallica_de_url_de_busca():
    url = "https://gallica.bnf.fr/SRU?operation=searchRetrieve&version=1.2&query=Christophori+Clavii"
    assert cli._extrair_consulta_gallica(url) == "Christophori Clavii"


def test_extrair_consulta_gallica_texto_direto():
    assert cli._extrair_consulta_gallica('gallica all "Clavius"') == 'gallica all "Clavius"'


def test_escolher_adapter_forcado_tem_prioridade():
    assert cli.escolher_adapter("https://qualquer-site.com", forcado="gallica") == "gallica"


def test_escolher_adapter_domino_desconhecido_levanta_erro():
    with pytest.raises(ValueError):
        cli.escolher_adapter("https://site-nao-cadastrado.com")


def test_main_dominio_desconhecido_mostra_mensagem_amigavel(capsys):
    codigo = cli.main(["https://site-nao-cadastrado.com"])
    assert codigo == 1
    assert "Não deu para continuar" in capsys.readouterr().out


def test_main_acao_humana_necessaria_mostra_aviso_amigavel(capsys, monkeypatch):
    monkeypatch.setattr(cli, "construir_adapter", lambda nome, url: AdapterFakePrecisaChave(url))

    codigo = cli.main(["https://grand-sud-medieval.fr/forum/viewtopic.php?f=14&t=1", "--adapter", "phpbb"])

    saida = capsys.readouterr().out
    assert codigo == 1
    assert "Preciso da sua ajuda" in saida
    assert "tal-site" in saida
    assert "cadastre-se" in saida


def test_main_acao_humana_necessaria_dispara_notificacao_windows(monkeypatch):
    chamadas = []
    monkeypatch.setattr(cli, "construir_adapter", lambda nome, url: AdapterFakePrecisaChave(url))
    monkeypatch.setattr(cli, "avisar_windows", lambda titulo, mensagem: chamadas.append((titulo, mensagem)))

    cli.main(["https://grand-sud-medieval.fr/forum/viewtopic.php?f=14&t=1", "--adapter", "phpbb"])

    assert len(chamadas) == 1
    titulo, mensagem = chamadas[0]
    assert "tal-site" in mensagem
    assert "cadastre-se" in mensagem


def test_construir_adapter_gallica():
    adapter = cli.construir_adapter("gallica", "Christophori Clavii")
    assert isinstance(adapter, GallicaAdapter)
    assert adapter.consulta == "Christophori Clavii"


def test_main_ponta_a_ponta_com_adapter_fake(tmp_path, monkeypatch):
    monkeypatch.setattr(cli, "construir_adapter", lambda nome, url: AdapterFake(url))
    monkeypatch.setattr("buscador.core.enriquecimento.analisar", lambda link: (
        ("vivo (pdf direto)", "verde") if link.endswith(".pdf") else ("requer login", "requer login")
    ))
    monkeypatch.setattr("buscador.core.enriquecimento.traduzir",
                         lambda texto, idioma_origem="auto", idioma_destino="pt": f"{texto} (PT)")

    destino = tmp_path / "saida.xlsx"
    codigo = cli.main([
        "https://grand-sud-medieval.fr/forum/viewtopic.php?f=14&t=1",
        "--adapter", "phpbb",
        "--saida", str(destino),
    ])

    assert codigo == 0
    assert destino.exists()

    ws = load_workbook(destino)["Links"]
    assert ws.max_row == 3  # cabecalho + 2 itens
    assert ws.cell(row=2, column=1).value == "Obra Teste"
    assert ws.cell(row=2, column=2).value == "Obra Teste (PT)"
    assert ws.cell(row=2, column=6).value == "pdf direto"
    assert ws.cell(row=3, column=6).value == "precisa login"


def test_main_usa_caminho_padrao_quando_sem_saida(tmp_path, monkeypatch):
    monkeypatch.setattr(cli, "construir_adapter", lambda nome, url: AdapterFake(url))
    monkeypatch.setattr("buscador.core.enriquecimento.analisar", lambda link: ("vivo (pdf direto)", "verde"))
    monkeypatch.setattr("buscador.core.enriquecimento.traduzir",
                         lambda texto, idioma_origem="auto", idioma_destino="pt": texto)
    monkeypatch.setattr(cli, "SAIDAS", tmp_path)

    codigo = cli.main(["https://grand-sud-medieval.fr/forum/viewtopic.php?f=14&t=1", "--adapter", "phpbb"])

    assert codigo == 0
    gerados = list(tmp_path.glob("grand-sud-medieval-fr_*.xlsx"))
    assert len(gerados) == 1
