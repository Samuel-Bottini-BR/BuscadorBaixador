# -*- coding: utf-8 -*-
import csv
import json
import time
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from buscador.adapters.base import Item
from buscador.core import enriquecimento_lote as modulo
from buscador.core.enriquecimento_lote import (
    CheckpointEnriquecimento,
    carregar_ou_criar_checkpoint,
    enriquecer_em_lotes,
    salvar_checkpoint,
)


def _itens(quantidade, idioma="fr"):
    return [
        Item(titulo_original=f"Livro {i}", link=f"https://exemplo.com/{i}", extra={"idioma_origem": idioma})
        for i in range(quantidade)
    ]


def _fake_enriquecer_item(item):
    item.status_link = "vivo (pdf direto)"
    item.tipo = "pdf-direto"
    item.titulo_pt = f"{item.titulo_original} (PT)"
    return item


def test_processa_em_lotes_gera_planilha_e_csv_por_lote(tmp_path, monkeypatch):
    monkeypatch.setattr(modulo, "enriquecer_item", _fake_enriquecer_item)
    monkeypatch.setattr(modulo, "traduzir", MagicMock())
    diretorio_job = tmp_path / "job"

    checkpoint = enriquecer_em_lotes(_itens(5), diretorio_job, lote_tamanho=2, workers=2)

    assert checkpoint.proximo_lote == 3
    assert (diretorio_job / "planilha_lote_0001.xlsx").exists()
    assert (diretorio_job / "planilha_lote_0002.xlsx").exists()
    assert (diretorio_job / "planilha_lote_0003.xlsx").exists()

    ws = load_workbook(diretorio_job / "planilha_lote_0001.xlsx")["Links"]
    assert ws.max_row == 3  # cabecalho + 2 itens
    assert ws.cell(row=2, column=2).value == "Livro 0 (PT)"  # coluna titulo_pt

    with open(diretorio_job / "itens_enriquecidos.csv", newline="", encoding="utf-8") as arquivo:
        linhas = list(csv.DictReader(arquivo))
    assert len(linhas) == 5
    assert linhas[0]["titulo_pt"] == "Livro 0 (PT)"
    assert linhas[0]["status_link"] == "vivo (pdf direto)"


def test_retomar_pula_lotes_ja_marcados_como_prontos(tmp_path, monkeypatch):
    chamadas = []

    def enriquecer_espiao(item):
        chamadas.append(item.titulo_original)
        return _fake_enriquecer_item(item)

    monkeypatch.setattr(modulo, "enriquecer_item", enriquecer_espiao)
    monkeypatch.setattr(modulo, "traduzir", MagicMock())
    diretorio_job = tmp_path / "job"

    # simula que o lote 1 (indice 0) ja foi processado numa execucao anterior
    checkpoint = carregar_ou_criar_checkpoint(diretorio_job / "checkpoint_enriquecimento.json",
                                               total_itens=5, lote_tamanho=2)
    checkpoint.proximo_lote = 1
    salvar_checkpoint(checkpoint, diretorio_job / "checkpoint_enriquecimento.json")

    checkpoint_final = enriquecer_em_lotes(_itens(5), diretorio_job, lote_tamanho=2, workers=2)

    assert checkpoint_final.proximo_lote == 3
    assert chamadas == ["Livro 2", "Livro 3", "Livro 4"]  # nao repetiu "Livro 0"/"Livro 1"
    assert not (diretorio_job / "planilha_lote_0001.xlsx").exists()  # nao foi tocado de novo
    assert (diretorio_job / "planilha_lote_0002.xlsx").exists()
    assert (diretorio_job / "planilha_lote_0003.xlsx").exists()


def test_pre_aquece_traducao_uma_vez_por_idioma_distinto(tmp_path, monkeypatch):
    monkeypatch.setattr(modulo, "enriquecer_item", _fake_enriquecer_item)
    traduzir_fake = MagicMock()
    monkeypatch.setattr(modulo, "traduzir", traduzir_fake)
    diretorio_job = tmp_path / "job"

    itens = _itens(2, idioma="fr") + _itens(2, idioma="la") + _itens(1, idioma="fr")
    enriquecer_em_lotes(itens, diretorio_job, lote_tamanho=10, workers=2)

    idiomas_aquecidos = {chamada.kwargs["idioma_origem"] for chamada in traduzir_fake.call_args_list}
    assert idiomas_aquecidos == {"fr", "la"}
    assert traduzir_fake.call_count == 2  # um por idioma distinto, nao um por item


def test_salvar_checkpoint_tenta_de_novo_quando_o_windows_nega_o_os_replace_por_um_instante(tmp_path, monkeypatch):
    # enquanto um leitor (o 'status' lendo o checkpoint) tem o arquivo aberto, o
    # os.replace do escritor falha com PermissionError [WinError 5] no Windows; sem
    # tentar de novo, isso mataria a coleta de horas por causa de um instante de azar
    caminho = tmp_path / "checkpoint_enriquecimento.json"
    salvar_checkpoint(CheckpointEnriquecimento(total_itens=100, lote_tamanho=20, proximo_lote=1), caminho)
    replace_original = modulo.os.replace
    tentativas = []

    def replace_negado_nas_duas_primeiras_vezes(origem, destino):
        tentativas.append(1)
        if len(tentativas) <= 2:
            raise PermissionError("[WinError 5] Acesso negado")
        return replace_original(origem, destino)
    monkeypatch.setattr(modulo.os, "replace", replace_negado_nas_duas_primeiras_vezes)
    esperas = []
    monkeypatch.setattr(time, "sleep", esperas.append)  # nao espera de verdade, mas registra

    salvar_checkpoint(CheckpointEnriquecimento(total_itens=100, lote_tamanho=20, proximo_lote=2), caminho)

    assert len(tentativas) == 3  # 2 negadas + 1 que funcionou
    assert len(esperas) == 2  # e esperou um pouco entre as tentativas (nao e um laco ocupado)
    assert json.loads(caminho.read_text(encoding="utf-8"))["proximo_lote"] == 2
    assert not caminho.with_suffix(".json.tmp").exists()


def test_salvar_checkpoint_relevanta_o_permission_error_se_o_windows_nega_sempre(tmp_path, monkeypatch):
    caminho = tmp_path / "checkpoint_enriquecimento.json"
    salvar_checkpoint(CheckpointEnriquecimento(total_itens=100, lote_tamanho=20, proximo_lote=1), caminho)
    tentativas = []

    def replace_sempre_negado(origem, destino):
        tentativas.append(1)
        raise PermissionError("[WinError 5] Acesso negado")
    monkeypatch.setattr(modulo.os, "replace", replace_sempre_negado)
    monkeypatch.setattr(time, "sleep", lambda segundos: None)

    with pytest.raises(PermissionError):
        salvar_checkpoint(CheckpointEnriquecimento(total_itens=100, lote_tamanho=20, proximo_lote=2), caminho)

    assert len(tentativas) == 10  # desiste depois de 10 tentativas, nao antes e nao pra sempre
    assert json.loads(caminho.read_text(encoding="utf-8"))["proximo_lote"] == 1  # o antigo segue valido


def test_lista_vazia_nao_quebra(tmp_path, monkeypatch):
    monkeypatch.setattr(modulo, "enriquecer_item", _fake_enriquecer_item)
    monkeypatch.setattr(modulo, "traduzir", MagicMock())

    checkpoint = enriquecer_em_lotes([], tmp_path / "job", lote_tamanho=10)

    assert checkpoint.proximo_lote == 0
    assert checkpoint.total_itens == 0
