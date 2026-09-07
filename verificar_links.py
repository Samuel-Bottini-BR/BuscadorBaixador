# -*- coding: utf-8 -*-
"""
Verificador de links da planilha do Garimpo (Instituto São Bento).

Testa cada link externo da aba "Links" (em paralelo) e PINTA A LINHA INTEIRA:
  🔴 VERMELHO  = quebrado (404/410/5xx/DNS/timeout)
  🔵 AZUL      = requer login (401/403/429 ou academia.edu, Drive, JSTOR, ResearchGate)
  🟢 VERDE     = vivo, sem login E com PDF baixável (URL .pdf, PDF no HTML, ou repositório conhecido)
  ⬜ BRANCO    = vivo, mas sem PDF detectável (linha sem cor)

Preenche a coluna 'status_link', cria as abas "Quebrados" e "Prontos p/ baixar",
e salva <nome>_VERIFICADO.xlsx (não sobrescreve o original). Não apaga nada.

Uso:
    python verificar_links.py                    (usa o .xlsx da pasta)
    python verificar_links.py "arquivo.xlsx"
Requisitos:  pip install requests openpyxl
"""
import sys, glob, os, re, concurrent.futures as cf

# Console do Windows às vezes não é UTF-8 (cp1252) e derruba o print dos emojis
# do resultado final — mesmo com a planilha já salva. Força UTF-8 na saída.
for _stream in (sys.stdout, sys.stderr):
    try: _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception: pass

try:
    import requests
except ImportError:
    print("Falta 'requests'. Rode:  pip install requests openpyxl"); sys.exit(1)
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

requests.packages.urllib3.disable_warnings()

TIMEOUT = 12
WORKERS = 40
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124 Safari/537.36"
LOGIN_HOSTS = ("academia.edu", "drive.google.com", "docs.google.com", "jstor.org", "researchgate.net")
# repositórios de texto integral: se estiver vivo, assume-se que há PDF (mesmo sem "ver" no HTML)
REPO_HOSTS = ("persee.fr", "archives-ouvertes.fr", "raco.cat", "revues.org", "openedition",
              "oapen.org", "unesdoc.unesco.org", "collectionscanada.gc.ca", "thesescanada",
              "archipel.uqam.ca", "theses.univ-lyon2.fr", "gallica.bnf.fr", "digidol.llgc.org.uk",
              "gipuzkoakultura.net", "nauticalarch.org", "bibnum.enc.sorbonne.fr")
PDF_RE = re.compile(r'\.pdf(["\'?#)\s]|$)', re.I)

COR = {
    "quebrado":     PatternFill("solid", fgColor="FFC7CE"),  # vermelho
    "requer login": PatternFill("solid", fgColor="BDD7EE"),  # azul
    "verde":        PatternFill("solid", fgColor="C6EFCE"),  # verde
}

def host_de(url):
    try: return url.split("/")[2].lower()
    except Exception: return ""

def tem_login(host):
    return any(h in host for h in LOGIN_HOSTS)

def eh_repo(host):
    return any(h in host for h in REPO_HOSTS)

def analisar(url):
    """Retorna (status_texto, categoria) onde categoria in {quebrado,requer login,verde,branco}."""
    host = host_de(url)
    if tem_login(host):
        return "requer login", "requer login"
    try:
        # 1) status + tipo de conteúdo (HEAD; se recusado, GET)
        r = None
        try:
            r = requests.head(url, timeout=TIMEOUT, allow_redirects=True,
                              headers={"User-Agent": UA}, verify=False)
            if r.status_code in (405, 501, 403):  # muitos servidores recusam HEAD
                r = None
        except requests.exceptions.SSLError:
            return "vivo (ssl)", "branco"
        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout):
            r = None

        if r is None:
            try:
                r = requests.get(url, timeout=TIMEOUT, allow_redirects=True,
                                 headers={"User-Agent": UA}, verify=False, stream=True)
            except requests.exceptions.SSLError:
                return "vivo (ssl)", "branco"
            except (requests.exceptions.ConnectionError, requests.exceptions.Timeout):
                return "quebrado", "quebrado"

        code = r.status_code
        ct = (r.headers.get("Content-Type") or "").lower()
        if code in (401, 403, 429):
            try: r.close()
            except Exception: pass
            return "requer login", "requer login"
        if code == 404 or code == 410 or code >= 500:
            try: r.close()
            except Exception: pass
            return "quebrado", "quebrado"
        if not (200 <= code < 400):
            try: r.close()
            except Exception: pass
            return "quebrado", "quebrado"

        # vivo. É PDF direto?
        if url.lower().split("?")[0].endswith(".pdf") or "application/pdf" in ct:
            try: r.close()
            except Exception: pass
            return "vivo (pdf direto)", "verde"

        # baixa um pedaço do HTML e procura link .pdf
        html = ""
        try:
            g = requests.get(url, timeout=TIMEOUT, allow_redirects=True,
                             headers={"User-Agent": UA}, verify=False, stream=True)
            if "application/pdf" in (g.headers.get("Content-Type") or "").lower():
                g.close(); return "vivo (pdf direto)", "verde"
            chunk = g.raw.read(600000, decode_content=True) or b""
            html = chunk.decode("utf-8", "ignore")
            g.close()
        except Exception:
            html = ""
        if PDF_RE.search(html):
            return "vivo (pdf na página)", "verde"
        # regra extra: repositório conhecido de texto integral
        if eh_repo(host):
            return "vivo (repositório)", "verde"
        return "vivo (sem pdf)", "branco"
    except Exception:
        return "quebrado", "quebrado"

def achar_arquivo():
    if len(sys.argv) > 1: return sys.argv[1]
    cand = [f for f in glob.glob("*.xlsx") if "VERIFICADO" not in f]
    cand.sort(key=lambda f: (0 if "MESTRE" in f.upper() else 1, f))
    if not cand:
        print("Nenhum .xlsx nesta pasta. Passe o caminho como argumento."); sys.exit(1)
    return cand[0]

def main():
    caminho = achar_arquivo()
    print("Lendo:", caminho)
    wb = load_workbook(caminho)
    ws = wb["Links"]
    header = [c.value for c in ws[1]]
    ci_link = (header.index("link") if "link" in header else header.index("link_encontrado"))
    ci_status = header.index("status_link")
    ci_titulo = header.index("titulo_pt") if "titulo_pt" in header else None
    ci_secao = header.index("secao") if "secao" in header else None
    ncols = len(header)

    linhas = list(ws.iter_rows(min_row=2))
    urls = {}
    for row in linhas:
        u = row[ci_link].value
        if u and "://" in str(u):
            urls.setdefault(str(u), None)
    print(f"{len(urls)} links únicos para verificar (de {len(linhas)} linhas). Aguarde alguns minutos...")

    feito = 0
    with cf.ThreadPoolExecutor(max_workers=WORKERS) as ex:
        futs = {ex.submit(analisar, u): u for u in urls}
        for fut in cf.as_completed(futs):
            urls[futs[fut]] = fut.result()
            feito += 1
            if feito % 250 == 0: print(f"  ...{feito}/{len(urls)}")

    from collections import Counter
    cont = Counter(); dom_falha = Counter()
    quebrados, verdes = [], []
    for row in linhas:
        u = row[ci_link].value
        res = urls.get(str(u)) if u else None
        if not res: continue
        texto, cat = res
        row[ci_status].value = texto
        cont[cat] += 1
        if cat in COR:  # pinta a linha inteira
            fill = COR[cat]
            for c in range(ncols):
                row[c].fill = fill
        tit = row[ci_titulo].value if ci_titulo is not None else ""
        sec = row[ci_secao].value if ci_secao is not None else ""
        if cat == "quebrado":
            quebrados.append([u, tit, sec]); dom_falha[host_de(str(u))] += 1
        elif cat == "verde":
            verdes.append([tit, sec, u])

    for nome in ("Quebrados", "Prontos para baixar"):
        if nome in wb.sheetnames: del wb[nome]
    wq = wb.create_sheet("Quebrados")
    wq.append(["link", "titulo_pt", "secao"])
    for q in quebrados: wq.append(q)
    wv = wb.create_sheet("Prontos para baixar")
    wv.append(["titulo_pt", "secao", "link"])
    for v in verdes: wv.append(v)
    for w in (wq, wv):
        for c in w[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
        w.column_dimensions["A"].width = 50; w.column_dimensions["B"].width = 30; w.column_dimensions["C"].width = 50
        w.freeze_panes = "A2"

    saida = os.path.splitext(caminho)[0] + "_VERIFICADO.xlsx"
    wb.save(saida)
    print("\n==== RESULTADO ====")
    print(f"  🟢 verde  (vivo + PDF baixável): {cont.get('verde',0)}")
    print(f"  ⬜ branco (só vivo, sem PDF)   : {cont.get('branco',0)}")
    print(f"  🔵 azul   (requer login)       : {cont.get('requer login',0)}")
    print(f"  🔴 vermelho (quebrado)         : {cont.get('quebrado',0)}")
    if dom_falha:
        print("\nDomínios que mais falharam:")
        for d, n in dom_falha.most_common(12): print(f"   {n:4d}  {d}")
    print("\nAbas criadas: 'Quebrados' e 'Prontos para baixar'.")
    print("Nada foi apagado. Arquivo salvo:", saida)

if __name__ == "__main__":
    main()
