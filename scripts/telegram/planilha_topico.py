"""Gera planilha .xlsx (Resumo + Arquivos) a partir do export do tdl de um tópico.

Uso: python planilha_topico.py <export.json> <saida.xlsx> <id_topico> "<nome do tópico>" [chat] ["<nome do grupo>"]
id_topico 0 = grupo/conversa sem tópicos. Sem [chat], usa o Refúgio Intelectual.
"""
import json
import os
import shutil
import sys
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

CHAT = 2136545743

GRUPOS = [
    ("PDF", {"pdf"}, ("application/pdf",)),
    ("EPUB/MOBI/AZW", {"epub", "mobi", "azw", "azw3", "fb2"}, ("application/epub",)),
    ("DJVU", {"djvu", "djv"}, ("image/vnd.djvu",)),
    ("DOC/DOCX/ODT/RTF/TXT", {"doc", "docx", "odt", "rtf", "txt"}, ("application/msword", "text/")),
    ("Compactado (ZIP/RAR/7Z)", {"zip", "rar", "7z", "tar", "gz"}, ("application/zip", "application/rar", "application/x-rar", "application/x-7z")),
    ("Vídeo", {"mp4", "mkv", "avi", "mov", "webm", "flv"}, ("video/",)),
    ("Áudio", {"mp3", "m4a", "ogg", "opus", "wav", "flac", "aac"}, ("audio/",)),
    ("Figurinha", {"webp", "tgs"}, ("image/webp", "application/x-tgsticker")),
    ("Imagem", {"jpg", "jpeg", "png", "gif", "tif", "tiff"}, ("image/",)),
]


def tipo_de(ext, mime):
    for nome, exts, _ in GRUPOS:
        if ext in exts:
            return nome
    for nome, _, mimes in GRUPOS:
        if any(mime.startswith(m) for m in mimes):
            return nome
    return "Sem extensão (provável PDF)" if mime == "application/octet-stream" else f"Outro ({mime or '?'})"


def info(raw):
    media = raw.get("Media") or {}
    doc = media.get("Document")
    if doc:
        nome = next((a["FileName"] for a in doc.get("Attributes") or [] if a.get("FileName")), "")
        return nome, doc.get("MimeType", ""), int(doc.get("Size") or 0)
    foto = media.get("Photo")
    if foto:
        tams = [s.get("Size") or max(s.get("Sizes") or [0]) for s in foto.get("Sizes") or []]
        return "", "image/jpeg", int(max(tams or [0]))
    return "", "", 0


def extensao(nome):
    base, ext = os.path.splitext(nome)
    ext = ext.lower().lstrip(".")
    # "Tomo.1", "Vol. 2" etc. não são extensões
    return ext if ext and ext.isalnum() and not ext.isdigit() and len(ext) <= 5 and " " not in ext else ""


def main(export, saida, topico, nome_topico, chat=CHAT, nome_grupo="Refúgio Intelectual"):
    dados = json.load(open(export, encoding="utf-8"))
    linhas = []
    for m in dados["messages"]:
        raw = m.get("raw") or {}
        nome, mime, tam = info(raw)
        nome = nome or m.get("file") or ""
        ext = extensao(nome) or ("jpg" if mime == "image/jpeg" and not nome else "")
        linhas.append({
            "id": m["id"],
            "data": datetime.fromtimestamp(m["date"]),
            "arquivo": nome,
            "ext": ext,
            "tipo": tipo_de(ext, mime),
            "mb": round(tam / 1024**2, 2),
            "bytes": tam,
            "legenda": (raw.get("Message") or "").replace("\n", " ").strip()[:500],
            "link": (f"https://t.me/c/{chat}/{topico}/{m['id']}" if topico
                     else f"https://t.me/c/{chat}/{m['id']}"),
        })
    linhas.sort(key=lambda r: r["id"])

    por_nome = defaultdict(list)
    for r in linhas:
        por_nome[r["arquivo"]].append(r)
    for nome, grupo in por_nome.items():
        if len(grupo) < 2:
            continue
        tamanhos = Counter(r["bytes"] for r in grupo)
        for r in grupo:
            r["obs"] = ("Repetido (mesmo nome e tamanho)" if tamanhos[r["bytes"]] > 1
                        else "Mesmo nome, arquivo diferente (volume?)")

    wb = Workbook()
    negrito = Font(bold=True)
    cab = PatternFill("solid", fgColor="DDE6F0")

    # Aba Resumo
    ws = wb.active
    ws.title = "Resumo"
    total = sum(r["bytes"] for r in linhas)
    ws.append([f"Tópico: {nome_topico} (grupo {nome_grupo}, tópico {topico})" if topico
               else f"Conversa: {nome_grupo}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Lista gerada em {datetime.now():%d/%m/%Y %H:%M} — nada foi baixado ainda"])
    ws.append([])
    ws.append(["Total de arquivos", len(linhas)])
    ws.append(["Tamanho total (GB)", round(total / 1024**3, 2)])
    ws.append(["Período das postagens", f"{linhas[0]['data']:%d/%m/%Y} a {linhas[-1]['data']:%d/%m/%Y}"])
    ws.append(["Repetidos (mesmo nome e tamanho)", sum(1 for r in linhas if r.get("obs", "").startswith("Repetido"))])
    ws.append(["Mesmo nome, arquivo diferente", sum(1 for r in linhas if r.get("obs", "").startswith("Mesmo nome"))])
    ws.append([])
    ws.append(["Tipo", "Arquivos", "Tamanho (GB)"])
    for c in ws[ws.max_row]:
        c.font, c.fill = negrito, cab
    por_tipo = defaultdict(lambda: [0, 0])
    for r in linhas:
        por_tipo[r["tipo"]][0] += 1
        por_tipo[r["tipo"]][1] += r["bytes"]
    for tipo, (n, b) in sorted(por_tipo.items(), key=lambda kv: -kv[1][1]):
        ws.append([tipo, n, round(b / 1024**3, 2)])
    for row in ws.iter_rows(min_row=4, max_row=8, max_col=1):
        row[0].font = negrito
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14

    # Aba Arquivos
    wa = wb.create_sheet("Arquivos")
    colunas = ["Nº mensagem", "Data", "Nome do arquivo", "Tipo", "Tamanho (MB)", "Observação", "Legenda", "Link"]
    wa.append(colunas)
    for c in wa[1]:
        c.font, c.fill = negrito, cab
    for r in linhas:
        wa.append([r["id"], r["data"], r["arquivo"], r["tipo"], r["mb"], r.get("obs", ""), r["legenda"], "abrir"])
        cel = wa.cell(row=wa.max_row, column=8)
        cel.hyperlink = r["link"]
        cel.font = Font(color="0563C1", underline="single")
        wa.cell(row=wa.max_row, column=2).number_format = "DD/MM/YYYY"
    larguras = [12, 12, 70, 24, 13, 36, 50, 8]
    for i, w in enumerate(larguras, 1):
        wa.column_dimensions[get_column_letter(i)].width = w
    wa.freeze_panes = "A2"
    wa.auto_filter.ref = wa.dimensions
    for row in wa.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].alignment = Alignment(wrap_text=False)

    wb.save(saida)

    print(f"Arquivos: {len(linhas)} | Total: {total / 1024**3:.2f} GB | "
          f"{linhas[0]['data']:%d/%m/%Y} a {linhas[-1]['data']:%d/%m/%Y}")
    for tipo, (n, b) in sorted(por_tipo.items(), key=lambda kv: -kv[1][1]):
        print(f"  {tipo:30s} {n:5d}  {b / 1024**3:7.2f} GB")
    obs = Counter(r.get("obs", "") for r in linhas if r.get("obs"))
    for k, v in obs.items():
        print(f"  {k}: {v}")
    print("Maiores:")
    for r in sorted(linhas, key=lambda r: -r["bytes"])[:5]:
        print(f"  {r['mb']:8.1f} MB  {r['arquivo']}")
    livre = shutil.disk_usage(os.path.dirname(os.path.abspath(saida))).free
    print(f"Espaço livre no disco da planilha: {livre / 1024**3:.0f} GB")
    print(f"Planilha: {saida}")


if __name__ == "__main__":
    extras = {}
    if len(sys.argv) > 5:
        extras["chat"] = int(sys.argv[5])
    if len(sys.argv) > 6:
        extras["nome_grupo"] = sys.argv[6]
    main(sys.argv[1], sys.argv[2], int(sys.argv[3]), sys.argv[4], **extras)
