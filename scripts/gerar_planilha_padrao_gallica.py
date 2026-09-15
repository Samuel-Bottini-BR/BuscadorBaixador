# -*- coding: utf-8 -*-
"""Gera a planilha no PADRAO APROVADO pelo Samuel em 15/09/2026: aba
"Indice" (categoria + quantidade + link clicavel) mais uma aba por
categoria (Titulo original / Traducao (PT) / Link).

Roda em cima da saida de categorizar_amostra_gallica.py. Ainda usa a
amostra pequena (24 itens) - para rodar nos 46.222 itens de verdade, falta:
(a) refinar os BALDES de categorizar_amostra_gallica.py (so 4 dos 8
pegaram itens na rodada de teste), (b) traduzir todo mundo, o que para
46 mil titulos via Google/MyMemory (rate limit real, ver HANDOFF.md) vai
precisar de lotes retomaveis, no estilo de core/enriquecimento_lote.py -
nao rodar tudo de uma vez numa chamada so.
"""
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

RAIZ = Path(__file__).resolve().parent.parent
ORIGEM = RAIZ / "saidas" / "amostra_categorizada_traduzida.json"
SAIDA = RAIZ / "saidas" / "gallica_livros_padrao.xlsx"

CABECALHO_FUNDO = "1F4E78"
CABECALHO_FONTE = Font(color="FFFFFF", bold=True)


def estilizar_cabecalho(ws, linha, num_colunas):
    for c in range(1, num_colunas + 1):
        cel = ws.cell(row=linha, column=c)
        cel.font = CABECALHO_FONTE
        cel.fill = PatternFill("solid", fgColor=CABECALHO_FUNDO)
        cel.alignment = Alignment(vertical="center")


def nome_de_aba(nome, usados):
    limpo = nome[:28]
    base = limpo
    i = 2
    while limpo in usados:
        limpo = f"{base[:25]}({i})"
        i += 1
    usados.add(limpo)
    return limpo


def main():
    dados = json.loads(ORIGEM.read_text(encoding="utf-8"))
    por_categoria = defaultdict(list)
    for item in dados:
        por_categoria[item["categoria_padronizada"]].append(item)

    wb = Workbook()
    indice = wb.active
    indice.title = "Indice"
    indice.append(["Categoria", "Quantidade", "Ir para a aba"])
    estilizar_cabecalho(indice, 1, 3)
    indice.column_dimensions["A"].width = 35
    indice.column_dimensions["B"].width = 15
    indice.column_dimensions["C"].width = 18

    usados = {"Indice"}
    linha = 2
    for categoria, itens in sorted(por_categoria.items(), key=lambda kv: -len(kv[1])):
        nome_aba = nome_de_aba(categoria, usados)
        indice.cell(row=linha, column=1, value=categoria)
        indice.cell(row=linha, column=2, value=len(itens))
        link = indice.cell(row=linha, column=3, value="abrir")
        link.hyperlink = f"#'{nome_aba}'!A1"
        link.font = Font(color="0563C1", underline="single")
        linha += 1

        ws = wb.create_sheet(nome_aba)
        ws.append(["Titulo original", "Traducao (PT)", "Link"])
        estilizar_cabecalho(ws, 1, 3)
        ws.freeze_panes = "A2"
        for item in itens:
            ws.append([item["titulo"], item["titulo_traduzido"], item["link"]])
        ws.column_dimensions["A"].width = 45
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 45
        ws.auto_filter.ref = f"A1:C{len(itens) + 1}"
        volta = ws.cell(row=1, column=5, value="< voltar ao indice")
        volta.hyperlink = "#'Indice'!A1"
        volta.font = Font(color="0563C1", underline="single")

    wb.save(SAIDA)
    print(f"Salvo: {SAIDA}")


if __name__ == "__main__":
    main()
